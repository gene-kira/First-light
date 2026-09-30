#!/usr/bin/env python3
"""Windows Universal Media-to-AI Bridge v2.0
Local-first Windows GUI for extracting useful metadata/text from images, video,
audio, documents and code, then optionally sending the extracted report to a
configured local AI endpoint. Requires optional external tools for some formats.
"""
from __future__ import annotations
import os, sys, json, mimetypes, platform, shutil, subprocess, threading, queue, traceback
from pathlib import Path
from datetime import datetime
import tkinter as tk
from tkinter import ttk, filedialog, messagebox

APP = "Windows Universal Media-to-AI Bridge"
VERSION = "2.0"
IMAGES = {".png",".jpg",".jpeg",".bmp",".gif",".webp",".tif",".tiff",".ico",".heic",".avif"}
VIDEOS = {".mp4",".mkv",".mov",".avi",".webm",".m4v",".mpeg",".mpg",".wmv",".flv",".ts",".mts",".m2ts",".3gp",".ogv"}
AUDIO = {".mp3",".wav",".flac",".m4a",".aac",".ogg",".opus",".wma",".aiff",".aif",".amr"}
TEXT = {".txt",".md",".rst",".csv",".tsv",".json",".xml",".yaml",".yml",".ini",".cfg",".log",".py",".js",".ts",".html",".css",".sql",".ps1",".bat",".cmd",".c",".cpp",".h",".java",".rs",".go",".toml",".srt",".vtt",".ass",".ssa"}
DOCS = {".pdf",".docx"}
LIMIT = 100_000

def stamp():
    return datetime.now().strftime("%Y-%m-%d %H:%M:%S")

def command(args, timeout=60):
    try:
        p = subprocess.run(args, capture_output=True, text=True, timeout=timeout, errors="replace",
                           creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0))
        return p.returncode, p.stdout.strip(), p.stderr.strip()
    except FileNotFoundError:
        return 127, "", f"Command not found: {args[0]}"
    except Exception as e:
        return 1, "", str(e)

def read_text(path):
    if path.stat().st_size > 20 * 1024 * 1024:
        raise ValueError("Text file exceeds 20 MB safety limit.")
    for enc in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
        try: return path.read_text(encoding=enc)
        except UnicodeDecodeError: pass
    return path.read_text(errors="replace")

def ffprobe(path):
    exe = shutil.which("ffprobe")
    if not exe: return None, "ffprobe not found; install FFmpeg and add its bin folder to PATH."
    rc, out, err = command([exe, "-v", "error", "-show_format", "-show_streams", "-of", "json", str(path)], 60)
    if rc: return None, err
    try: return json.loads(out), None
    except Exception as e: return None, str(e)

def analyze(filename):
    p = Path(filename).expanduser().resolve()
    if not p.is_file(): raise FileNotFoundError(str(p))
    if p.stat().st_size > 1024*1024*1024: raise ValueError("File exceeds 1 GB safety limit.")
    ext = p.suffix.lower()
    mime, _ = mimetypes.guess_type(str(p))
    r = {"app":APP,"version":VERSION,"created_at":stamp(),"source_file":str(p),
         "file_name":p.name,"extension":ext,"mime_type":mime or "application/octet-stream",
         "size_bytes":p.stat().st_size,"size_mb":round(p.stat().st_size/1048576,3),
         "modified_at":datetime.fromtimestamp(p.stat().st_mtime).isoformat(timespec="seconds"),
         "platform":platform.platform()}
    m = {"kind":"unknown"}
    if ext in IMAGES:
        m["kind"]="image"
        try:
            from PIL import Image
            with Image.open(p) as im:
                m.update(format=im.format,width=im.width,height=im.height,mode=im.mode,frames=getattr(im,"n_frames",1))
            try:
                import pytesseract
                m["ocr_text"]=pytesseract.image_to_string(Image.open(p))[:LIMIT]
            except ImportError: m["ocr_note"]="Optional OCR requires pytesseract and Tesseract OCR installed."
            except Exception as e: m["ocr_note"]=f"OCR unavailable: {e}"
        except ImportError: m["note"]="Install Pillow for image inspection: python -m pip install pillow"
        except Exception as e: m["image_error"]=str(e)
        r["summary"]="Image metadata and optional OCR. Visual interpretation requires a vision-capable model."
    elif ext in VIDEOS:
        m["kind"]="video"
        data, err = ffprobe(p)
        if data:
            fmt=data.get("format",{})
            m["format"]={k:fmt.get(k) for k in ("format_name","duration","size","bit_rate","tags")}
            m["streams"]=[{k:s.get(k) for k in ("index","codec_type","codec_name","width","height","avg_frame_rate","sample_rate","channels","duration","bit_rate","tags") if k in s} for s in data.get("streams",[])]
            try: duration=float(fmt.get("duration") or 0)
            except (ValueError,TypeError): duration=0
            exe=shutil.which("ffmpeg")
            m["frames_sampled"]=[]
            if exe and duration>0:
                folder=p.parent/(p.stem+"_ai_frames"); folder.mkdir(exist_ok=True)
                for i in range(5):
                    sec=duration*(i+1)/6
                    dest=folder/f"frame_{i+1:02d}.jpg"
                    rc,_,e=command([exe,"-y","-ss",f"{sec:.3f}","-i",str(p),"-frames:v","1","-vf","scale=960:-1",str(dest)],45)
                    if rc==0 and dest.exists(): m["frames_sampled"].append({"time_seconds":round(sec,3),"path":str(dest)})
                    elif e: m.setdefault("frame_errors",[]).append(e[:300])
            elif not exe: m["frame_note"]="Install FFmpeg to extract sample frames."
        else: m["error"]=err
        r["summary"]="Video metadata/streams and optional sampled frames via FFmpeg."
    elif ext in AUDIO:
        m["kind"]="audio"
        data,err=ffprobe(p)
        if data:
            m["format"]=data.get("format",{})
            m["streams"]=[{k:s.get(k) for k in ("codec_type","codec_name","sample_rate","channels","duration","bit_rate","tags") if k in s} for s in data.get("streams",[])]
        else: m["metadata_note"]=err
        whisper=shutil.which("whisper")
        if whisper:
            rc,out,e=command([whisper,str(p),"--output_format","txt","--output_dir",str(p.parent)],1800)
            txt=p.with_suffix(".txt")
            if rc==0 and txt.exists():
                try: m["transcript"]=read_text(txt)[:LIMIT]; m["transcript_file"]=str(txt)
                except Exception as ex: m["transcription_error"]=str(ex)
            else: m["transcription_note"]=(e or out or "Whisper did not create a transcript.")[:1500]
        else: m["transcription_note"]="Optional: install local Whisper and ensure the whisper command is on PATH."
        r["summary"]="Audio metadata and optional local Whisper transcription."
    elif ext in DOCS:
        m["kind"]="document"
        try:
            if ext==".pdf":
                from pypdf import PdfReader
                reader=PdfReader(str(p)); m["pages"]=len(reader.pages)
                chunks=[]
                for i,page in enumerate(reader.pages):
                    chunks.append(f"[Page {i+1}]\n{page.extract_text() or ''}")
                    if sum(map(len,chunks))>=LIMIT: break
                m["text"]="\n\n".join(chunks)[:LIMIT]
            else:
                from docx import Document
                doc=Document(str(p)); m["text"]="\n".join(x.text for x in doc.paragraphs)[:LIMIT]
                m["tables"]=[[[c.text for c in row.cells] for row in t.rows] for t in doc.tables[:20]]
        except ImportError: m["note"]="Install pypdf or python-docx for document extraction."
        except Exception as e: m["extraction_error"]=str(e)
        r["summary"]="Extracted document text. Scanned PDFs may require OCR."
    elif ext in TEXT:
        m={"kind":"text","text":read_text(p)[:LIMIT]}
        r["summary"]="Text/code/subtitle content extracted."
    else:
        r["summary"]="Basic file metadata only; no decoder is configured for this extension."
        m["note"]="Unsupported or unknown format."
    r["media"]=m
    return r

def post_json(url,payload,timeout=180):
    import urllib.request
    req=urllib.request.Request(url,data=json.dumps(payload).encode(),headers={"Content-Type":"application/json"})
    with urllib.request.urlopen(req,timeout=timeout) as response:
        return json.loads(response.read().decode("utf-8",errors="replace"))

class App:
    def __init__(self,root,initial):
        self.root=root; root.title(f"{APP} v{VERSION}"); root.geometry("1050x720"); root.minsize(800,540)
        self.files=[]; self.reports={}; self.busy=False; self.events=queue.Queue()
        self.outdir=Path.home()/"MediaAI_Bridge"; self.outdir.mkdir(exist_ok=True)
        self.provider=tk.StringVar(value="Ollama"); self.url=tk.StringVar(value="http://127.0.0.1:11434")
        self.model=tk.StringVar(value="llava"); self.status=tk.StringVar(value="Ready.")
        self.ui(); self.add_paths(initial); root.after(100,self.poll)

    def ui(self):
        ttk.Label(self.root,text=APP,font=("Segoe UI",15,"bold"),padding=10).pack(anchor="w")
        nb=ttk.Notebook(self.root); nb.pack(fill="both",expand=True,padx=10,pady=5)
        work=ttk.Frame(nb,padding=8); ai=ttk.Frame(nb,padding=12); win=ttk.Frame(nb,padding=12); diag=ttk.Frame(nb,padding=8)
        for f,title in ((work,"Media Workspace"),(ai,"AI Connection"),(win,"Windows Integration"),(diag,"Diagnostics / Log")): nb.add(f,text=title)
        bar=ttk.Frame(work); bar.pack(fill="x",pady=(0,8))
        for label,fn in (("Add Files",self.pick),("Add Folder",self.folder),("Remove",self.remove),("Analyze Selected",self.selected),("Analyze All",self.all),("Save Report",self.save),("Ask AI",self.ask)):
            ttk.Button(bar,text=label,command=fn).pack(side="left",padx=2)
        pan=ttk.Panedwindow(work,orient="horizontal"); pan.pack(fill="both",expand=True)
        left=ttk.Frame(pan); right=ttk.Frame(pan); pan.add(left,weight=1); pan.add(right,weight=2)
        ttk.Label(left,text="Files").pack(anchor="w")
        self.lst=tk.Listbox(left,selectmode="extended"); self.lst.pack(fill="both",expand=True); self.lst.bind("<<ListboxSelect>>",self.display)
        ttk.Label(right,text="Extracted report").pack(anchor="w")
        self.view=tk.Text(right,wrap="word",font=("Consolas",9)); self.view.pack(fill="both",expand=True)
        ttk.Label(ai,text="Provider").grid(row=0,column=0,sticky="w",pady=5)
        ttk.Combobox(ai,textvariable=self.provider,values=["Ollama","LM Studio","OpenAI-compatible"],state="readonly").grid(row=0,column=1,sticky="w")
        ttk.Label(ai,text="Base URL").grid(row=1,column=0,sticky="w",pady=5); ttk.Entry(ai,textvariable=self.url,width=55).grid(row=1,column=1,sticky="ew")
        ttk.Label(ai,text="Model").grid(row=2,column=0,sticky="w",pady=5); ttk.Entry(ai,textvariable=self.model,width=35).grid(row=2,column=1,sticky="w")
        ttk.Button(ai,text="Test Connection",command=self.test).grid(row=3,column=1,sticky="w",pady=8)
        ttk.Label(ai,text="Only the extracted report text/metadata is sent. Original media bytes are not uploaded. Review reports first.").grid(row=4,column=0,columnspan=2,sticky="w")
        ai.columnconfigure(1,weight=1)
        ttk.Label(win,text="Windows Explorer integration",font=("Segoe UI",12,"bold")).pack(anchor="w",pady=(0,8))
        ttk.Label(win,text="Install a per-user Send To shortcut. In Explorer, right-click a file → Send to → this app. No registry changes are made.").pack(anchor="w",fill="x",pady=5)
        ttk.Button(win,text="Install Explorer Send To Shortcut",command=self.install).pack(anchor="w",pady=4)
        ttk.Button(win,text="Remove Explorer Send To Shortcut",command=self.uninstall).pack(anchor="w",pady=4)
        ttk.Button(win,text="Open Report Folder",command=lambda:self.open(self.outdir)).pack(anchor="w",pady=4)
        ttk.Label(win,text=f"Report folder: {self.outdir}").pack(anchor="w",pady=8)
        ttk.Label(win,text="Optional standalone EXE:\npython -m pip install pyinstaller\npython -m PyInstaller --noconsole --onefile --name MediaAIBridge Windows_Universal_Media_to_AI_Bridge.py").pack(anchor="w",pady=12)
        row=ttk.Frame(diag); row.pack(fill="x")
        ttk.Button(row,text="Run Diagnostics",command=self.diagnostics).pack(side="left")
        self.logbox=tk.Text(diag,wrap="word",font=("Consolas",9)); self.logbox.pack(fill="both",expand=True,pady=5)
        ttk.Label(self.root,textvariable=self.status,relief="sunken",anchor="w",padding=5).pack(fill="x",side="bottom")
        self.log("Application started.")

    def log(self,msg):
        line=f"[{stamp()}] {msg}"
        self.logbox.insert("end",line+"\n"); self.logbox.see("end")

    def add_paths(self,paths):
        n=0
        for x in paths or []:
            p=Path(x)
            if p.is_file() and str(p.resolve()) not in self.files:
                s=str(p.resolve()); self.files.append(s); self.lst.insert("end",s); n+=1
        if n: self.status.set(f"Added {n} file(s).")

    def pick(self): self.add_paths(filedialog.askopenfilenames(title="Select media and documents"))
    def folder(self):
        d=filedialog.askdirectory()
        if d: self.add_paths([str(p) for p in Path(d).rglob("*") if p.is_file() and p.suffix.lower() in IMAGES|VIDEOS|AUDIO|TEXT|DOCS][:1000])
    def chosen(self): return [self.files[i] for i in self.lst.curselection()]
    def remove(self):
        for i in reversed(self.lst.curselection()):
            self.reports.pop(self.files[i],None); self.files.pop(i); self.lst.delete(i)
    def display(self,_=None):
        p=(self.chosen() or [None])[0]
        if p:
            self.view.delete("1.0","end"); self.view.insert("1.0",json.dumps(self.reports[p],indent=2,ensure_ascii=False) if p in self.reports else p+"\nNot analyzed yet.")
    def selected(self):
        ps=self.chosen()
        if not ps: messagebox.showinfo("Select files","Select one or more files first."); return
        self.start(ps)
    def all(self):
        if not self.files: self.pick()
        if self.files: self.start(list(self.files))
    def start(self,paths):
        if self.busy: messagebox.showinfo("Busy","An analysis batch is running."); return
        self.busy=True
        def work():
            for p in paths:
                try:
                    self.events.put(("status",f"Analyzing {Path(p).name}..."))
                    r=analyze(p); self.reports[p]=r
                    safe="".join(c if c.isalnum() or c in "-_." else "_" for c in Path(p).stem)[:100]
                    dest=self.outdir/f"{safe}_media_report.json"; dest.write_text(json.dumps(r,indent=2,ensure_ascii=False),encoding="utf-8")
                    self.events.put(("report",p,r)); self.events.put(("status",f"Analyzed {Path(p).name}; saved {dest}"))
                except Exception as e: self.events.put(("error",f"{p}: {e}"))
            self.events.put(("done",))
        threading.Thread(target=work,daemon=True).start()
    def poll(self):
        try:
            while True:
                e=self.events.get_nowait()
                if e[0]=="status": self.status.set(e[1]); self.log(e[1])
                elif e[0]=="report":
                    p,r=e[1:]; self.view.delete("1.0","end"); self.view.insert("1.0",json.dumps(r,indent=2,ensure_ascii=False))
                    try: self.lst.selection_clear(0,"end"); self.lst.selection_set(self.files.index(p))
                    except ValueError: pass
                elif e[0]=="error": self.log("ERROR: "+e[1])
                elif e[0]=="ai": self.view.insert("end","\n\n--- AI RESPONSE ---\n"+e[1]); self.view.see("end"); self.status.set("AI response received.")
                elif e[0]=="done": self.busy=False; self.status.set("Analysis complete.")
        except queue.Empty: pass
        self.root.after(100,self.poll)
    def save(self):
        ps=self.chosen()
        if not ps or ps[0] not in self.reports: messagebox.showinfo("No report","Analyze and select a file first."); return
        dest=filedialog.asksaveasfilename(defaultextension=".json",initialfile=Path(ps[0]).stem+"_report.json",filetypes=[("JSON","*.json")])
        if dest: Path(dest).write_text(json.dumps(self.reports[ps[0]],indent=2,ensure_ascii=False),encoding="utf-8")
    def request_ai(self,prompt):
        base=self.url.get().strip().rstrip("/"); model=self.model.get().strip()
        if not base or not model: raise ValueError("Enter a base URL and model.")
        if self.provider.get()=="Ollama":
            url=base if base.endswith("/api/generate") else base+"/api/generate"
            d=post_json(url,{"model":model,"prompt":prompt,"stream":False})
            return d.get("response",json.dumps(d))
        url=base if base.endswith("/v1/chat/completions") else base+"/v1/chat/completions"
        d=post_json(url,{"model":model,"messages":[{"role":"user","content":prompt}],"stream":False})
        return d.get("choices",[{}])[0].get("message",{}).get("content",json.dumps(d))
    def ask(self):
        ps=self.chosen()
        if not ps or ps[0] not in self.reports: messagebox.showinfo("No report","Analyze and select a file first."); return
        prompt="Analyze this extracted media report. State what can be concluded, uncertainties, and useful next steps. Do not claim to have seen original media if only metadata/text is provided.\n\n"+json.dumps(self.reports[ps[0]],ensure_ascii=False)[:LIMIT]
        self.status.set("Sending report to AI...")
        threading.Thread(target=self._ai_worker,args=(prompt,),daemon=True).start()
    def test(self): threading.Thread(target=self._ai_worker,args=("Reply exactly: Connection successful.",),daemon=True).start()
    def _ai_worker(self,prompt):
        try: self.events.put(("ai",self.request_ai(prompt)))
        except Exception as e: self.events.put(("ai",f"AI request failed: {e}"))
    def install(self):
        if os.name!="nt": messagebox.showwarning("Windows only","Explorer integration works only on Windows."); return
        folder=Path(os.environ.get("APPDATA",str(Path.home())))/"Microsoft/Windows/SendTo"; folder.mkdir(parents=True,exist_ok=True)
        target=folder/"Windows Universal Media-to-AI Bridge.cmd"
        if getattr(sys,"frozen",False): cmd=f'"{sys.executable}" %*'
        else: cmd=f'"{sys.executable}" "{Path(__file__).resolve()}" %*'
        target.write_text("@echo off\r\n"+cmd+"\r\n",encoding="utf-8")
        messagebox.showinfo("Installed",f"Created:\n{target}")
    def uninstall(self):
        target=Path(os.environ.get("APPDATA",str(Path.home())))/"Microsoft/Windows/SendTo/Windows Universal Media-to-AI Bridge.cmd"
        try: target.unlink(missing_ok=True); self.status.set("Explorer shortcut removed.")
        except Exception as e: messagebox.showerror("Error",str(e))
    def open(self,p):
        try:
            if os.name=="nt": os.startfile(str(p))
            elif sys.platform=="darwin": subprocess.Popen(["open",str(p)])
            else: subprocess.Popen(["xdg-open",str(p)])
        except Exception as e: messagebox.showerror("Open folder failed",str(e))
    def diagnostics(self):
        lines=[f"{APP} {VERSION}",f"Python: {sys.version}",f"Platform: {platform.platform()}",
               f"FFmpeg: {shutil.which('ffmpeg') or 'not found'}",f"ffprobe: {shutil.which('ffprobe') or 'not found'}",
               f"Whisper: {shutil.which('whisper') or 'not found'}"]
        for mod in ("PIL","pytesseract","pypdf","docx"):
            try: __import__(mod); lines.append(f"{mod}: available")
            except Exception: lines.append(f"{mod}: not installed")
        self.log("\n".join(lines)); messagebox.showinfo("Diagnostics","\n".join(lines))

def main():
    root=tk.Tk(); App(root,[p for p in sys.argv[1:] if os.path.isfile(p)]); root.mainloop()

if __name__=="__main__":
    try: main()
    except Exception:
        traceback.print_exc()
        try: input("Application error. Press Enter to close...")
        except Exception: pass
